"""Generate the knowledge extraction e2e corpus and its ground truth.

Run from the repository root:
uv run --no-project --with reportlab --with python-docx --with openpyxl --with mongoengine --with boto3 --with httpx \
  python e2e/fixtures/knowledge-extraction/generate_corpus.py

Invoices state their supplier, number and date below the page-header zone, since MinerU drops lines it reads as a
running header. The ground truth lists, per file, the records an extraction must find:

- `ground_truth.json`: hardware lines per invoice (with `borderline` for items a reader may or may not count);
- `ground_truth_contracts.json`: notice period, liability cap and governing law per supplier contract;
- `ground_truth_slas.json`: response and resolution time per priority, and the availability penalty, per SLA.
"""

import json
import random
import shutil
import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).parent))
from common import (  # noqa: E402
    BULK_FILES,
    CORPUS,
    GROUND_TRUTH,
    GROUND_TRUTH_CONTRACTS,
    GROUND_TRUTH_SLAS,
)

# Enough documents that one answer cannot extract them all: the sample, the facets and the parts are exercised.
HOOLI_FILES = 150
# Above the default full-read limit of 500,000 characters, so the document is read in part.
LONG_DOCUMENT_CHARACTERS = 600_000

TRUTH: dict[str, list[dict[str, object]]] = {}
CONTRACTS: dict[str, dict[str, str]] = {}
SLAS: dict[str, dict[str, object]] = {}

HW, SW, SV = "Hardware", "Software", "Service"


def record(collection: str, path: str, hardware: list[tuple[str, float]], supplier: str, year: int) -> None:
    TRUTH[f"{collection}/{path}"] = [
        {"item": item, "amount": amount, "supplier": supplier, "year": year} for item, amount in hardware
    ]


def write_pdf(target: Path, pages: list[list[str]]) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(target), pagesize=A4)
    for lines in pages:
        y = 700
        for line in lines:
            pdf.drawString(60, y, line)
            y -= 18
        pdf.showPage()
    pdf.save()


def write_text(collection: str, path: str, text: str) -> None:
    target = CORPUS / collection / path
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def invoice_lines(number: str, supplier: str, date: str, items: list[tuple[str, str, float]]) -> list[str]:
    lines = [f"Invoice number: {number}", f"Supplier: {supplier}", f"Invoice date: {date}", "", "Positions:"]
    lines += [f"{index}. {category}: {item} - CHF {amount:.2f}" for index, (category, item, amount) in enumerate(items, 1)]
    total = sum(amount for _, _, amount in items)
    return [*lines, "", f"Total: CHF {total:.2f}"]


def pdf_invoice(collection, path, number, supplier, date, items) -> None:
    write_pdf(CORPUS / collection / path, [invoice_lines(number, supplier, date, items)])
    record(collection, path, [(i, a) for c, i, a in items if c == HW], supplier, int(date[:4]))


def markdown_invoice(collection, path, number, supplier, date, items) -> None:
    rows = "\n".join(f"| {n} | {c} | {i} | {a:.2f} |" for n, (c, i, a) in enumerate(items, 1))
    write_text(
        collection,
        path,
        f"# Invoice {number}\n\nSupplier: {supplier}\n\nInvoice date: {date}\n\n"
        f"| # | Category | Item | Amount (CHF) |\n|---|---|---|---|\n{rows}\n",
    )
    record(collection, path, [(i, a) for c, i, a in items if c == HW], supplier, int(date[:4]))


def docx_invoice(collection, path, number, supplier, date, items) -> None:
    document = Document()
    document.add_heading(f"Invoice {number}", 1)
    document.add_paragraph(f"Supplier: {supplier}")
    document.add_paragraph(f"Invoice date: {date}")
    table = document.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = "Category", "Item", "Amount CHF"
    for category, item, amount in items:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = category, item, f"{amount:.2f}"
    target = CORPUS / collection / path
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))
    record(collection, path, [(i, a) for c, i, a in items if c == HW], supplier, int(date[:4]))


def xlsx_invoice(collection, path, number, supplier, date, items) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row in (["Invoice number", number], ["Supplier", supplier], ["Invoice date", date], [], ["Category", "Item", "Amount CHF"]):
        sheet.append(row)
    for category, item, amount in items:
        sheet.append([category, item, amount])
    target = CORPUS / collection / path
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(target))
    record(collection, path, [(i, a) for c, i, a in items if c == HW], supplier, int(date[:4]))


def long_invoice() -> None:
    """Forty pages of services with one hardware line on the last page: windowing must not lose it."""
    path = "invoices/2025/q4/acme-2025-long.pdf"
    pages = [invoice_lines("ACME-2025-090", "ACME AG", "2025-11-30", [(SV, "Support hours October", 1200.0)])]
    for page in range(2, 40):
        pages.append([f"Service log page {page}"] + [f"Ticket {page}-{i}: remote support session, no material used." for i in range(1, 30)])
    pages.append(["Additional positions:", "Hardware: Spare power supply PSU-750 - CHF 310.00", "End of invoice ACME-2025-090"])
    write_pdf(CORPUS / "finance" / path, pages)
    record("finance", path, [("Spare power supply PSU-750", 310.0)], "ACME AG", 2025)


def acme_and_neighbours() -> None:
    """The corpus the first e2e cases (E1 to E24, U1 to U10) score against; unchanged so they stay comparable."""
    pdf_invoice("finance", "invoices/2025/q1/acme-2025-001.pdf", "ACME-2025-001", "ACME AG", "2025-02-14", [
        (HW, "Server rack R100", 1200.0), (HW, "Network switch 48 ports", 890.0), (HW, "Rack PDU", 310.0),
        (HW, "SSD 2 TB", 240.0), (HW, "Patch cables (50)", 95.0), (SV, "Installation", 500.0),
    ])
    docx_invoice("finance", "invoices/2025/q2/acme-2025-002.docx", "ACME-2025-002", "ACME AG", "2025-05-03", [
        (HW, "Laptop T14", 1450.0), (HW, "Docking station", 220.0), (SW, "Office licences (10)", 990.0),
    ])
    pdf_invoice("finance", "invoices/2025/q3/inv-77.pdf", "77", "ACME AG", "2025-08-21", [
        (HW, "Monitor 27 inch", 380.0), (SV, "Delivery", 40.0),
    ])
    pdf_invoice("finance", "invoices/2025/q4/acme-2025-003.pdf", "ACME-2025-003", "ACME AG", "2025-10-10", [
        (SV, "Consulting days (3)", 4200.0), (SW, "Backup software renewal", 600.0),
    ])
    long_invoice()
    pdf_invoice("finance", "invoices/2025/q3/globex-2025-010.pdf", "GLOBEX-2025-010", "Globex GmbH", "2025-07-01", [
        (HW, "Firewall appliance", 2300.0), (HW, "Access point", 180.0),
    ])
    xlsx_invoice("finance", "invoices/2025/q4/initech-2025-020.xlsx", "INITECH-2025-020", "Initech SA", "2025-12-01", [
        (SV, "Cloud hosting Q4", 1500.0), (HW, "UPS 1500 VA", 420.0),
    ])
    pdf_invoice("finance", "invoices/2024/acme-2024-050.pdf", "ACME-2024-050", "ACME AG", "2024-09-09", [
        (HW, "Monitor 24 inch", 250.0), (HW, "Keyboard", 45.0),
    ])
    markdown_invoice("finance", "contracts/acme-hardware-frame.md", "FRAME-01", "ACME AG", "2025-01-01", [
        (HW, "Framework hardware budget", 50000.0),
    ])
    for index in range(1, BULK_FILES + 1):
        markdown_invoice("finance", f"invoices/2025/bulk/acme-b-{index:03d}.md", f"ACME-B-{index:03d}", "ACME AG",
                         f"2025-06-{(index % 28) + 1:02d}", [
            (HW, f"USB-C hub HB{index}", 35.0 + index), (HW, f"Headset HS{index}", 60.0 + index), (SV, "Shipping", 15.0),
        ])
    write_text("legal", "contracts/msa-acme.md", "# Master services agreement with ACME AG\n\nCovers hardware maintenance 2025-2027.\n")
    pdf_invoice("hr", "invoices/2025/q1/acme-2025-001.pdf", "ACME-2025-901", "ACME AG", "2025-03-03", [
        (HW, "Ergonomic chair", 640.0), (SV, "Assembly", 80.0),
    ])


def hooli_bulk() -> None:
    """150 invoices of one supplier over three years and four quarters: too many for one answer."""
    for index in range(HOOLI_FILES):
        year, quarter = 2023 + index % 3, index % 4 + 1
        markdown_invoice(
            "finance",
            f"invoices/{year}/hooli/q{quarter}/hooli-{year}-{index:03d}.md",
            f"HOOLI-{year}-{index:03d}",
            "Hooli AG",
            f"{year}-{quarter * 3:02d}-{(index % 27) + 1:02d}",
            [(HW, f"Thin client TC{index}", 300.0 + index), (HW, f"Keyboard KB{index}", 40.0 + index % 10), (SV, "Setup", 50.0)],
        )


def vandelay_long() -> None:
    """An invoice longer than the full-read limit, its hardware lines in the middle and at the end."""
    path = "invoices/2025/q2/vandelay-2025-001.md"
    filler = random.Random(1888)
    lines = ["# Invoice VAN-2025-001", "", "Supplier: Vandelay Industries", "", "Invoice date: 2025-04-30", "", "## Service log", ""]
    log_line = "Ticket {n}: on-site support visit, {minutes} minutes, no material used, closed by technician {tech}."
    size = sum(len(line) + 1 for line in lines)
    middle_written = False
    n = 0
    while size < LONG_DOCUMENT_CHARACTERS:
        n += 1
        line = log_line.format(n=n, minutes=filler.randint(15, 240), tech=filler.choice("ABCDEFGH"))
        if not middle_written and size > LONG_DOCUMENT_CHARACTERS // 2:
            line = "Hardware: Workstation W-900 - CHF 2400.00"
            middle_written = True
        lines.append(line)
        size += len(line) + 1
    lines += ["", "## Additional positions", "", "Hardware: Spare SSD 4 TB - CHF 380.00", "", "End of invoice VAN-2025-001"]
    write_text("finance", path, "\n".join(lines) + "\n")
    record("finance", path, [("Workstation W-900", 2400.0), ("Spare SSD 4 TB", 380.0)], "Vandelay Industries", 2025)


def initrode() -> None:
    """Invoices whose hardware lines never say "hardware", with two items a reader may or may not count."""
    items = [
        [("Laptop X1 Carbon", 1900.0, False), ("27-inch monitor", 410.0, False), ("Consulting (2 days)", 2800.0, None)],
        [("Wireless mouse", 35.0, False), ("Printer toner cartridges", 120.0, True), ("Project management", 900.0, None)],
        [("Extended laptop warranty (3 years)", 250.0, True), ("Barcode scanner", 180.0, False)],
    ]
    for index, lines in enumerate(items, 1):
        path = f"invoices/2025/q3/initrode-2025-0{index}.pdf"
        number = f"INITRODE-2025-0{index}"
        text = [f"Invoice number: {number}", "Supplier: Initrode AG", "Invoice date: 2025-09-1{index}", "", "Positions:"]
        text += [f"{n}. {item} - CHF {amount:.2f}" for n, (item, amount, _) in enumerate(lines, 1)]
        write_pdf(CORPUS / "finance" / path, [text])
        TRUTH[f"finance/{path}"] = [
            {"item": item, "amount": amount, "supplier": "Initrode AG", "year": 2025, "borderline": borderline}
            for item, amount, borderline in lines
            if borderline is not None
        ]


CONTRACT_TEMPLATE = """# Supply agreement with {supplier}

Agreement number: {number}

## 1. Subject

{supplier} supplies IT hardware and the related services to the customer.

## 2. Term and termination

{termination}

## 3. Liability

{liability}

## 4. Governing law and jurisdiction

{law}
"""


def contracts() -> None:
    """Supplier contracts with the clauses a contracts profile extracts; one is longer than the full-read limit."""
    terms = [
        ("ACME AG", "SUP-2024-01", "3 months", "Either party may terminate this agreement with three (3) months' notice to the end of a calendar quarter.", "CHF 500,000", "Liability of each party is limited to CHF 500,000 per contract year.", "Swiss law, courts of Zurich", "This agreement is governed by Swiss law. Place of jurisdiction is Zurich."),
        ("Globex GmbH", "SUP-2024-02", "6 months", "The agreement may be terminated by either party giving six (6) months' written notice.", "twice the annual fees", "Total liability is capped at twice the annual fees paid under this agreement.", "German law, courts of Munich", "German law applies. Exclusive place of jurisdiction is Munich."),
        ("Initech SA", "SUP-2025-03", "90 days", "Termination requires ninety (90) days' notice in writing.", "CHF 1,000,000", "The supplier's liability shall not exceed CHF 1,000,000 in aggregate.", "Swiss law, courts of Geneva", "Swiss substantive law applies; the courts of Geneva have jurisdiction."),
        ("Hooli AG", "SUP-2025-04", "1 month", "This agreement can be terminated with one (1) month's notice to the end of any month.", "the fees of the last 12 months", "Liability is limited to the fees paid in the twelve months before the claim.", "Swiss law, courts of Bern", "Swiss law governs this agreement; place of jurisdiction is Bern."),
    ]
    for index, (supplier, number, notice, termination, cap, liability, law_summary, law) in enumerate(terms):
        path = f"contracts/supply-{supplier.split()[0].lower()}-{number.lower()}.md"
        write_text("legal", path, CONTRACT_TEMPLATE.format(supplier=supplier, number=number, termination=termination, liability=liability, law=law))
        CONTRACTS[f"legal/{path}"] = {"supplier": supplier, "notice_period": notice, "liability_cap": cap, "governing_law": law_summary}
    long_contract()


def long_contract() -> None:
    """An outsourcing contract above the full-read limit, its termination clause in the middle."""
    path = "contracts/outsourcing-vandelay-sup-2023-09.md"
    filler = random.Random(1999)
    schedule = "Schedule {n}: service description item {n}, scope as agreed in the statement of work, reviewed quarterly by {team}."
    lines = ["# Outsourcing agreement with Vandelay Industries", "", "Agreement number: SUP-2023-09", "", "## Schedules", ""]
    size, n, written = sum(len(line) + 1 for line in lines), 0, False
    while size < LONG_DOCUMENT_CHARACTERS:
        n += 1
        line = schedule.format(n=n, team=filler.choice(["operations", "security", "finance", "legal"]))
        if not written and size > LONG_DOCUMENT_CHARACTERS // 2:
            line = "Termination: either party may terminate this agreement with twelve (12) months' notice."
            written = True
        lines.append(line)
        size += len(line) + 1
    lines += ["", "## Governing law", "", "Swiss law applies; the courts of Lucerne have jurisdiction."]
    write_text("legal", path, "\n".join(lines) + "\n")
    CONTRACTS[f"legal/{path}"] = {"supplier": "Vandelay Industries", "notice_period": "12 months", "liability_cap": "", "governing_law": "Swiss law, courts of Lucerne"}


def slas() -> None:
    """Service level agreements with a priority table; two of them define a penalty for missed availability."""
    agreements = [
        ("ACME AG", "99.9%", {"P1": ("30 minutes", "4 hours"), "P2": ("2 hours", "1 business day"), "P3": ("1 business day", "5 business days")}, "5% of the monthly fee per 0.1% below the target"),
        ("Globex GmbH", "99.5%", {"P1": ("1 hour", "8 hours"), "P2": ("4 hours", "2 business days")}, None),
        ("Initech SA", "99.95%", {"P1": ("15 minutes", "2 hours"), "P2": ("1 hour", "8 hours"), "P3": ("4 hours", "3 business days"), "P4": ("2 business days", "10 business days")}, "a credit of 10% of the monthly fee"),
        ("Hooli AG", "99.0%", {"P1": ("2 hours", "1 business day")}, None),
    ]
    for supplier, availability, priorities, penalty in agreements:
        path = f"slas/sla-{supplier.split()[0].lower()}.md"
        rows = "\n".join(f"| {p} | {response} | {resolution} |" for p, (response, resolution) in priorities.items())
        penalty_text = f"If availability falls below the target, the supplier owes {penalty}." if penalty else "No penalty applies."
        write_text(
            "legal",
            path,
            f"# Service level agreement with {supplier}\n\n## Availability\n\nTarget availability: {availability} per month.\n\n"
            f"{penalty_text}\n\n## Response and resolution times\n\n| Priority | Response | Resolution |\n|---|---|---|\n{rows}\n",
        )
        SLAS[f"legal/{path}"] = {
            "supplier": supplier,
            "availability": availability,
            "priorities": {p: {"response": r, "resolution": s} for p, (r, s) in priorities.items()},
            "penalty": penalty,
        }


def main() -> None:
    shutil.rmtree(CORPUS, ignore_errors=True)
    acme_and_neighbours()
    hooli_bulk()
    vandelay_long()
    initrode()
    contracts()
    slas()
    GROUND_TRUTH.write_text(json.dumps(TRUTH, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    GROUND_TRUTH_CONTRACTS.write_text(json.dumps(CONTRACTS, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    GROUND_TRUTH_SLAS.write_text(json.dumps(SLAS, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    files = sum(1 for path in CORPUS.rglob("*") if path.is_file())
    print(f"{files} files; ground truth for {len(TRUTH)} invoices, {len(CONTRACTS)} contracts, {len(SLAS)} SLAs")


if __name__ == "__main__":
    main()
