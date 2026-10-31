"""The user's own file space over the code sandbox: their home only, as their own account, with sandbox refusals passed
on and every path kept inside the home."""

import io
import json
from typing import Any
from unittest.mock import AsyncMock, MagicMock, call, patch

import pytest
import pytest_asyncio
from asgi_lifespan import LifespanManager
from httpx import ASGITransport, AsyncClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from swiss_ai_hub.core.generative_ai import ExtractedDocument, ParseCacheBucket
from swiss_ai_hub.core.infrastructure import OpenTerminalError
from swiss_ai_hub.core.testing.auth_utils import TEST_USER_OID, TestAuthHandler, fake_user

from swiss_ai_hub.api.routes.user_knowledge import user_knowledge_service
from swiss_ai_hub.api.routes.user_knowledge.user_knowledge_controller import UserKnowledgeController
from swiss_ai_hub.api.runners.api_test_runner import ApiTestRunner

pytestmark = pytest.mark.beanie

ENDPOINT = "/api/v1/active/user-knowledge"
THREAD = "65f1c0ffee00000000000001"


@pytest_asyncio.fixture(scope="module", loop_scope="session")
async def client(beanie_test_database):
    runner = ApiTestRunner()
    controller = UserKnowledgeController(auth=TestAuthHandler())
    (
        controller.list_user_files()
        .get_user_file_content()
        .preview_user_file()
        .upload_user_file()
        .create_user_folder()
        .move_user_file()
        .delete_user_file()
    )
    runner.mount(controller)
    app = runner.create_app()
    async with LifespanManager(app) as lifespan:
        async with AsyncClient(transport=ASGITransport(app=lifespan.app), base_url="http://test") as http:
            yield http


@pytest.fixture
def sandbox() -> Any:
    fake = MagicMock()
    fake.list_files = AsyncMock(
        return_value={
            "dir": "/home/owui1234",
            "entries": [
                {"name": "report.pdf", "type": "file", "size": 1200, "modified": 1.0},
                {"name": ".cache", "type": "directory", "size": 4096, "modified": 1.0},
                {"name": "conversations", "type": "directory", "size": 4096, "modified": 2.0},
            ],
        }
    )
    fake.view = AsyncMock(return_value=(b"%PDF-1.7", "application/pdf"))
    fake.upload = AsyncMock(return_value={})
    fake.mkdir = AsyncMock(return_value={})
    fake.move = AsyncMock(return_value={})
    fake.delete = AsyncMock(return_value={})
    with (
        patch.object(
            user_knowledge_service.OpenWebuiAccountEntity, "openwebui_id_of", return_value="owui-1234"
        ) as id_of,
        patch.object(user_knowledge_service, "OpenTerminalClient", return_value=fake) as created,
    ):
        fake.id_of, fake.created = id_of, created
        yield fake


async def test_the_listing_is_the_user_s_own_home_without_dotfiles(client: AsyncClient, sandbox: Any) -> None:
    response = await client.get(ENDPOINT)

    assert response.status_code == 200
    sandbox.id_of.assert_called_once_with(TEST_USER_OID)
    sandbox.created.assert_called_once_with("owui-1234")
    assert [(entry["name"], entry["kind"]) for entry in response.json()["entries"]] == [
        ("conversations", "folder"),
        ("report.pdf", "file"),
    ]


async def test_conversation_folders_carry_the_chat_title(client: AsyncClient, sandbox: Any) -> None:
    sandbox.list_files.return_value = {"entries": [{"name": THREAD, "type": "directory", "modified": 1.0}]}
    with patch.object(
        user_knowledge_service.UserKnowledgeService, "_conversation_titles", return_value={THREAD: "Q1 sales"}
    ) as titles:
        response = await client.get(ENDPOINT, params={"folder": "conversations"})

    titles.assert_called_once()
    assert response.json()["entries"][0]["conversation_title"] == "Q1 sales"


async def test_a_conversation_folder_carries_its_own_chat_title(client: AsyncClient, sandbox: Any) -> None:
    sandbox.list_files.return_value = {"entries": [{"name": "brief.md", "type": "file", "size": 9, "modified": 1.0}]}
    with patch.object(
        user_knowledge_service.UserKnowledgeService, "_conversation_titles", return_value={THREAD: "Q1 sales"}
    ):
        response = await client.get(ENDPOINT, params={"folder": f"conversations/{THREAD}"})

    assert response.json()["folder_title"] == "Q1 sales"


async def test_an_own_conversation_without_files_yet_is_empty_and_titled(client: AsyncClient, sandbox: Any) -> None:
    sandbox.list_files.side_effect = OpenTerminalError("404: Directory not found", 404)
    with patch.object(
        user_knowledge_service.UserKnowledgeService, "_conversation_titles", return_value={THREAD: "Q1 sales"}
    ):
        response = await client.get(ENDPOINT, params={"folder": f"conversations/{THREAD}"})

    assert response.status_code == 200
    assert (response.json()["entries"], response.json()["folder_title"]) == ([], "Q1 sales")


async def test_a_missing_folder_that_is_not_an_own_conversation_stays_missing(
    client: AsyncClient, sandbox: Any
) -> None:
    sandbox.list_files.side_effect = OpenTerminalError("404: Directory not found", 404)
    with patch.object(user_knowledge_service.UserKnowledgeService, "_conversation_titles", return_value={}):
        response = await client.get(ENDPOINT, params={"folder": f"conversations/{THREAD}"})

    assert response.status_code == 404


async def test_a_download_returns_the_bytes_as_an_attachment(client: AsyncClient, sandbox: Any) -> None:
    response = await client.get(f"{ENDPOINT}/content", params={"path": "reports/Q1 report.pdf", "download": True})

    assert response.content == b"%PDF-1.7"
    assert response.headers["content-disposition"] == "attachment; filename*=UTF-8''Q1%20report.pdf"
    assert response.headers["x-content-type-options"] == "nosniff"
    sandbox.view.assert_awaited_once_with("reports/Q1 report.pdf")


async def test_a_pdf_is_shown_in_place(client: AsyncClient, sandbox: Any) -> None:
    response = await client.get(f"{ENDPOINT}/content", params={"path": "report.pdf"})

    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith("inline;")


@pytest.mark.parametrize(
    "content_type", ["text/html", "image/svg+xml", "application/xhtml+xml", "text/html; charset=utf-8"]
)
async def test_a_file_a_browser_could_run_script_from_is_never_shown_in_place(
    client: AsyncClient, sandbox: Any, content_type: str
) -> None:
    sandbox.view.return_value = (b"<script>alert(1)</script>", content_type)

    response = await client.get(f"{ENDPOINT}/content", params={"path": "page.html"})

    assert response.headers["content-type"] == "application/octet-stream"
    assert response.headers["content-disposition"].startswith("attachment;")
    assert response.headers["content-security-policy"] == "sandbox; default-src 'none'"


def _extracted(markdown: str) -> ExtractedDocument:
    return ExtractedDocument(
        title="Report",
        content=markdown,
        content_type="",
        source_filename="report.docx",
        document_parser="MarkItDownLoader",
    )


async def test_an_office_document_is_previewed_as_the_markdown_an_agent_reads(
    client: AsyncClient, sandbox: Any
) -> None:
    docx = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    sandbox.view.return_value = (b"PK-docx", f"{docx}; charset=binary")
    with patch.object(
        user_knowledge_service.DocumentExtractor,
        "extract_from_bytes",
        AsyncMock(return_value=_extracted("# Report\n\n| a | b |")),
    ) as extract:
        response = await client.get(f"{ENDPOINT}/preview", params={"path": "reports/Q1.DOCX"})

    assert response.status_code == 200
    assert response.json() == {"path": "reports/Q1.DOCX", "markdown": "# Report\n\n| a | b |", "truncated": False}
    sandbox.view.assert_awaited_once_with("reports/Q1.DOCX")
    extract.assert_awaited_once_with(content=b"PK-docx", filename="Q1.DOCX", content_type=docx)


def _deck_with_a_picture() -> bytes:
    picture = io.BytesIO()
    Image.new("RGB", (64, 64), "red").save(picture, "PNG")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Quarterly figures"
    slide.shapes.add_picture(picture, Inches(1), Inches(1))
    content = io.BytesIO()
    deck.save(content)
    return content.getvalue()


async def test_a_real_document_is_previewed_without_its_images(client: AsyncClient, sandbox: Any) -> None:
    pptx = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    sandbox.view.return_value = (_deck_with_a_picture(), pptx)
    with (
        patch.object(ParseCacheBucket, "read", return_value=None),
        patch.object(ParseCacheBucket, "write"),
    ):
        response = await client.get(f"{ENDPOINT}/preview", params={"path": "deck.pptx"})

    assert response.status_code == 200
    markdown = response.json()["markdown"]
    assert "# Quarterly figures" in markdown
    assert "[image: image.png]" in markdown
    assert "data:image" not in markdown


async def test_a_long_preview_is_cut_at_a_whole_line_and_says_so(client: AsyncClient, sandbox: Any) -> None:
    limit = user_knowledge_service.UserKnowledgeService.MAX_PREVIEW_CHARS
    line = "x" * 99 + "\n"
    with patch.object(
        user_knowledge_service.DocumentExtractor,
        "extract_from_bytes",
        AsyncMock(return_value=_extracted(line * (limit // len(line) + 1))),
    ):
        response = await client.get(f"{ENDPOINT}/preview", params={"path": "big.docx"})

    assert response.status_code == 200
    assert response.json()["markdown"] == line * (limit // len(line))
    assert response.json()["truncated"] is True


async def test_a_table_the_cut_lands_in_is_closed_after_its_last_whole_row(client: AsyncClient, sandbox: Any) -> None:
    limit = user_knowledge_service.UserKnowledgeService.MAX_PREVIEW_CHARS
    rows = "".join(f"| {index:08d} |\n" for index in range(limit // 13 + 1))
    with patch.object(
        user_knowledge_service.DocumentExtractor,
        "extract_from_bytes",
        AsyncMock(return_value=_extracted(f"<table>| id |\n| --- |\n{rows}</table>")),
    ):
        response = await client.get(f"{ENDPOINT}/preview", params={"path": "big.xlsx"})

    markdown = response.json()["markdown"]
    assert markdown.endswith(" |\n</table>")
    assert len(markdown) <= limit + len("</table>")


async def test_a_file_too_large_to_preview_is_refused_before_it_is_converted(client: AsyncClient, sandbox: Any) -> None:
    limit = user_knowledge_service.UserKnowledgeService.MAX_PREVIEW_BYTES
    sandbox.view.return_value = (b"x" * (limit + 1), "application/vnd.ms-excel")
    with patch.object(user_knowledge_service.DocumentExtractor, "extract_from_bytes", AsyncMock()) as extract:
        response = await client.get(f"{ENDPOINT}/preview", params={"path": "huge.xls"})

    assert response.status_code == 413
    extract.assert_not_awaited()


@pytest.mark.parametrize("path", ["report.pdf", "notes.txt", "page.html", "mail.eml", "archive"])
async def test_a_file_without_a_markdown_preview_is_refused_before_it_is_read(
    client: AsyncClient, sandbox: Any, path: str
) -> None:
    response = await client.get(f"{ENDPOINT}/preview", params={"path": path})

    assert response.status_code == 415
    sandbox.view.assert_not_awaited()


async def test_an_upload_lands_in_the_chosen_folder(client: AsyncClient, sandbox: Any) -> None:
    response = await client.post(f"{ENDPOINT}/files", params={"folder": "notes"}, files={"file": ("a.txt", b"hi")})

    assert response.json() == {"path": "notes/a.txt"}
    sandbox.upload.assert_awaited_once_with("notes", "a.txt", b"hi")


@pytest.mark.parametrize(
    ("method", "url", "kwargs"),
    [
        ("get", ENDPOINT, {"params": {"folder": "/etc"}}),
        ("get", f"{ENDPOINT}/content", {"params": {"path": "../../proc/1/environ"}}),
        ("post", f"{ENDPOINT}/folders", {"json": {"path": "~/../other"}}),
        ("post", f"{ENDPOINT}/move", {"json": {"source": "a.txt", "destination": "/tmp/a.txt"}}),
        ("delete", ENDPOINT, {"params": {"path": "."}}),
        ("get", f"{ENDPOINT}/content", {"params": {"path": ".ssh/id_rsa"}}),
        ("delete", ENDPOINT, {"params": {"path": f"conversations/{THREAD}/.attached_files.json"}}),
        ("post", f"{ENDPOINT}/files", {"params": {"folder": "."}, "files": {"file": (".bashrc", b"x")}}),
    ],
)
async def test_a_path_outside_the_home_or_the_top_itself_is_refused(
    client: AsyncClient, sandbox: Any, method: str, url: str, kwargs: dict
) -> None:
    response = await getattr(client, method)(url, **kwargs)

    assert response.status_code == 400
    sandbox.view.assert_not_awaited()
    sandbox.upload.assert_not_awaited()
    sandbox.move.assert_not_awaited()
    sandbox.delete.assert_not_awaited()


async def test_an_upload_over_the_limit_is_refused_before_it_reaches_the_sandbox(
    client: AsyncClient, sandbox: Any
) -> None:
    with patch.object(user_knowledge_service.UserKnowledgeService, "MAX_UPLOAD_BYTES", 3):
        response = await client.post(f"{ENDPOINT}/files", files={"file": ("a.txt", b"four")})

    assert response.status_code == 413
    sandbox.upload.assert_not_awaited()


async def test_a_rename_moves_within_the_folder(client: AsyncClient, sandbox: Any) -> None:
    response = await client.post(f"{ENDPOINT}/move", json={"source": "notes/a.txt", "destination": "notes/b.txt"})

    assert response.json() == {"path": "notes/b.txt"}
    sandbox.move.assert_awaited_once_with("notes/a.txt", "notes/b.txt")


async def test_a_user_without_a_chat_account_has_no_file_space_yet(client: AsyncClient, sandbox: Any) -> None:
    sandbox.id_of.return_value = None

    response = await client.get(ENDPOINT)

    assert response.status_code == 404
    sandbox.list_files.assert_not_awaited()


async def test_a_sandbox_refusal_is_passed_on_and_a_failure_is_a_bad_gateway(client: AsyncClient, sandbox: Any) -> None:
    sandbox.delete.side_effect = OpenTerminalError("404: Path not found", 404)
    missing = await client.delete(ENDPOINT, params={"path": "gone.txt"})
    sandbox.list_files.side_effect = OpenTerminalError("500: boom", 500)
    broken = await client.get(ENDPOINT)

    assert (missing.status_code, missing.json()["detail"]) == (404, "404: Path not found")
    assert (broken.status_code, broken.json()["detail"]) == (502, "The code sandbox did not respond.")


FILE_ID = "0f8fad5b-d9cb-469f-a165-70867728950e"
OTHER_FILE_ID = "7c9e6679-7425-40de-944b-e07fc1f90ae7"


def _record(sandbox: Any, placed: dict[str, str] | None) -> None:
    """A fresh conversation folder and its record of placed attachments, or none yet."""
    missing = OpenTerminalError("404: File not found", 404)
    sandbox.view = AsyncMock(
        side_effect=missing if placed is None else None, return_value=(json.dumps(placed).encode(), "text/plain")
    )
    sandbox.write_file = AsyncMock(return_value={})
    sandbox.list_files = AsyncMock(return_value={"entries": []})


def _read(file_id: str, filename: str) -> bytes:
    return f"{file_id}/{filename}".encode()


async def test_a_turn_s_attachments_are_placed_in_its_conversation_folder_together(sandbox: Any) -> None:
    _record(sandbox, None)

    placed = await user_knowledge_service.UserKnowledgeService.place_attachments(
        fake_user(), THREAD, {FILE_ID: "report.pdf", OTHER_FILE_ID: "notes.txt"}, _read
    )

    assert sandbox.upload.await_args_list == [
        call(f"conversations/{THREAD}", "report.pdf", f"{FILE_ID}/report.pdf".encode()),
        call(f"conversations/{THREAD}", "notes.txt", f"{OTHER_FILE_ID}/notes.txt".encode()),
    ]
    sandbox.write_file.assert_awaited_once()
    assert [entry.path for entry in placed] == [
        f"conversations/{THREAD}/report.pdf",
        f"conversations/{THREAD}/notes.txt",
    ]


async def test_an_attachment_already_placed_is_not_placed_again_on_a_later_turn(sandbox: Any) -> None:
    _record(sandbox, {FILE_ID: "report.pdf"})

    await user_knowledge_service.UserKnowledgeService.place_attachments(
        fake_user(), THREAD, {FILE_ID: "report.pdf", OTHER_FILE_ID: "notes.txt"}, _read
    )

    sandbox.upload.assert_awaited_once_with(
        f"conversations/{THREAD}", "notes.txt", f"{OTHER_FILE_ID}/notes.txt".encode()
    )


async def test_a_user_without_a_file_space_gets_no_attachment_placed(sandbox: Any) -> None:
    sandbox.id_of.return_value = None

    placed = await user_knowledge_service.UserKnowledgeService.place_attachments(
        fake_user(), THREAD, {FILE_ID: "a.pdf"}, _read
    )

    assert placed == []
    sandbox.upload.assert_not_awaited()
