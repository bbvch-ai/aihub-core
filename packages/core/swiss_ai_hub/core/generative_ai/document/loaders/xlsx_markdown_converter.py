import datetime
import io
from collections.abc import Iterable, Iterator, Sequence
from typing import TYPE_CHECKING

from openpyxl import load_workbook

from swiss_ai_hub.core.generative_ai.document.tables.markdown_table import markdown_cell, markdown_row

if TYPE_CHECKING:
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

EXCEL_SIGNIFICANT_DIGITS = 15


class XlsxMarkdownConverter:
    """
    An xlsx workbook as markdown: a `## <sheet>` heading and one table per sheet, the first non-empty row as its header.

    Rows are streamed from openpyxl's read-only reader straight into table lines. MarkItDown renders each sheet to HTML
    and parses it back into a BeautifulSoup tree at about 1.6 KB per cell: a 60,000-row sheet took ~950 MB and 40 s that
    way, enough to get the API killed for memory, against ~20 MB and 4 s here. Values come out as entered rather than
    through pandas, which wrote 42.31 as 4.231000e+01, 3 as 3.0 and empty cells as NaN.
    """

    @staticmethod
    def convert(file_bytes: bytes) -> str:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        try:
            return "\n\n".join(XlsxMarkdownConverter._sheet(sheet) for sheet in workbook.worksheets)
        finally:
            workbook.close()

    @staticmethod
    def _sheet(sheet: "ReadOnlyWorksheet") -> str:
        """The read-only reader trusts the size a sheet declares, and files not written by Excel can declare A1:A1 for
        a full sheet, so the declared size is dropped and every stored row read, as pandas does. Every row is padded to
        the widest, since `parse_markdown_table` drops a row whose cell count differs from its header's."""
        sheet.reset_dimensions()
        heading = f"## {sheet.title}"
        cells = list(XlsxMarkdownConverter._trimmed_rows(sheet.iter_rows(values_only=True)))
        if not cells:
            return heading
        width = max(len(row) for row in cells)
        header, *body = cells
        return "\n".join(
            [
                heading,
                markdown_row(XlsxMarkdownConverter._padded(header, width)),
                markdown_row(["---"] * width),
                *(markdown_row(XlsxMarkdownConverter._padded(row, width)) for row in body),
            ]
        )

    @staticmethod
    def _trimmed_rows(rows: Iterable[Sequence[object]]) -> Iterator[list[str]]:
        """Each row's cells without the empty ones at its end. Empty rows before the first and after the last filled one
        are dropped; those in between stay, since they separate blocks such as two tables on one sheet."""
        empty_rows = 0
        started = False
        for values in rows:
            cells = [markdown_cell(XlsxMarkdownConverter._text(value)) for value in values]
            while cells and not cells[-1]:
                cells.pop()
            if not cells:
                if started:
                    empty_rows += 1
                continue
            yield from ([] for _ in range(empty_rows))
            empty_rows = 0
            started = True
            yield cells

    @staticmethod
    def _padded(cells: list[str], width: int) -> list[str]:
        return cells + [""] * (width - len(cells))

    @staticmethod
    def _text(value: object) -> str:
        """A value as the sheet shows it: a float to the 15 digits Excel keeps, which drops binary noise such as
        0.30000000000000004, booleans as Excel writes them, and a date cell without the midnight time it carries."""
        match value:
            case None:
                return ""
            case bool():
                return "TRUE" if value else "FALSE"
            case float():
                return f"{value:.{EXCEL_SIGNIFICANT_DIGITS}g}"
            case datetime.datetime():
                moment = XlsxMarkdownConverter._whole_second(value)
                return moment.date().isoformat() if moment.time() == datetime.time() else moment.isoformat(sep=" ")
            case datetime.time():
                moment = XlsxMarkdownConverter._whole_second(datetime.datetime.combine(datetime.date.min, value))
                return moment.time().isoformat()
            case datetime.date():
                return value.isoformat()
            case _:
                return str(value)

    @staticmethod
    def _whole_second(moment: datetime.datetime) -> datetime.datetime:
        """Excel stores a time as a fraction of a day, so 15:30:18 reads back as 15:30:17.999; within a millisecond of
        a whole second, that second is what was entered."""
        rounded = (moment + datetime.timedelta(microseconds=500_000)).replace(microsecond=0)
        return rounded if abs(rounded - moment) <= datetime.timedelta(milliseconds=1) else moment
